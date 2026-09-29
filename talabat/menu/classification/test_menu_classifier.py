"""
test_menu_classifier.py
Pytest suite for menu_classifier.py

Covers:
  - AsyncTokenBucket rate limiting
  - build_system_prompt output
  - Checkpoint save/load/roundtrip
  - load_input Excel parsing
  - API response parsing (valid terms, invalid terms, casing, edge cases)
  - Retry on RateLimitError
  - write_excel output structure
  - write_json output structure
  - Integration: process_batch end-to-end

Run:
  cd talabat/menu/classification
  pytest test_menu_classifier.py -v

Or from project root:
  pytest talabat/menu/classification/test_menu_classifier.py -v
"""

import asyncio
import json
import sys
import time
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

import httpx
import openpyxl
import pytest
import pytest_asyncio  # noqa: F401  (required for async tests)

# ── Import the module under test ──────────────────────────────────────────────
sys.path.insert(0, str(Path(__file__).parent))
import menu_classifier as mc


# ══════════════════════════════════════════════════════════════════════════════
# Helpers
# ══════════════════════════════════════════════════════════════════════════════

def _make_xlsx(path: Path, data_rows: list[list]) -> None:
    """Create a minimal 4-column test Excel file (header + data rows)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Item Name", "Category", "Description", "existing_std_term"])
    for row in data_rows:
        ws.append(row)
    wb.save(path)


def _mock_openai(response_text: str):
    """Return an AsyncMock OpenAI client whose create() yields response_text."""
    client = AsyncMock()
    resp = MagicMock()
    resp.choices = [MagicMock()]
    resp.choices[0].message.content = response_text
    client.chat.completions.create = AsyncMock(return_value=resp)
    return client


def _make_classify_fn(response_text: str, valid_terms: set):
    """Wire up a classify_fn with a mocked OpenAI client."""
    bucket = mc.AsyncTokenBucket(rpm=100_000)   # high RPM = no delays
    sem    = asyncio.Semaphore(10)
    client = _mock_openai(response_text)
    return mc.make_classifier(client, bucket, sem, "sys-prompt", valid_terms)


# ══════════════════════════════════════════════════════════════════════════════
# 1. AsyncTokenBucket
# ══════════════════════════════════════════════════════════════════════════════

@pytest.mark.asyncio
async def test_bucket_first_acquire_is_instant():
    """Full bucket → first acquire should not block."""
    bucket = mc.AsyncTokenBucket(rpm=600)
    t0 = time.monotonic()
    await bucket.acquire()
    assert time.monotonic() - t0 < 0.1, "First acquire should be instant"


@pytest.mark.asyncio
async def test_bucket_high_rpm_no_meaningful_delay():
    """10 rapid acquires at 100K RPM should finish in < 200 ms."""
    bucket = mc.AsyncTokenBucket(rpm=100_000)
    t0 = time.monotonic()
    for _ in range(10):
        await bucket.acquire()
    assert time.monotonic() - t0 < 0.2


@pytest.mark.asyncio
async def test_bucket_waits_when_tokens_exhausted():
    """With RPM=60 and tokens forced to 0.5, acquire must sleep ~0.5 s."""
    bucket = mc.AsyncTokenBucket(rpm=60)   # 1 token/s
    bucket._tokens = 0.5
    bucket._last   = time.monotonic()
    t0 = time.monotonic()
    await bucket.acquire()
    elapsed = time.monotonic() - t0
    assert 0.3 < elapsed < 2.0, f"Expected ~0.5s wait; got {elapsed:.3f}s"


@pytest.mark.asyncio
async def test_bucket_refills_over_time():
    """After 1 second a 60-RPM bucket should have ~1 new token."""
    bucket = mc.AsyncTokenBucket(rpm=60)
    bucket._tokens = 0.0
    bucket._last   = time.monotonic() - 1.0   # pretend 1 s passed
    # Should not raise; token was refilled by time passage
    t0 = time.monotonic()
    await bucket.acquire()
    assert time.monotonic() - t0 < 0.1


# ══════════════════════════════════════════════════════════════════════════════
# 2. build_system_prompt
# ══════════════════════════════════════════════════════════════════════════════

def test_prompt_contains_all_terms():
    terms = ["burger", "pizza", "shawarma", "kunafa"]
    prompt = mc.build_system_prompt(terms)
    for t in terms:
        assert t in prompt, f"Term '{t}' missing from system prompt"


def test_prompt_output_format_instructions():
    prompt = mc.build_system_prompt(["burger"])
    assert "INDEX|standard_term" in prompt
    assert "Others" in prompt


def test_prompt_tiered_steps_present():
    """All three classification steps must appear in the prompt."""
    prompt = mc.build_system_prompt(["burger", "juice"])
    assert "STEP 1" in prompt, "STEP 1 (item name) missing"
    assert "STEP 2" in prompt, "STEP 2 (category) missing"
    assert "STEP 3" in prompt, "STEP 3 (description) missing"
    assert "FALLBACK" in prompt, "FALLBACK clause missing"


def test_prompt_step_priority_order():
    """Steps must appear in the correct order: 1 before 2 before 3 before FALLBACK."""
    prompt = mc.build_system_prompt(["burger"])
    pos1 = prompt.index("STEP 1")
    pos2 = prompt.index("STEP 2")
    pos3 = prompt.index("STEP 3")
    posF = prompt.index("FALLBACK")
    assert pos1 < pos2 < pos3 < posF, "Step order is wrong in system prompt"


def test_prompt_step1_name_signal():
    """STEP 1 must reference the item name as the primary signal."""
    prompt = mc.build_system_prompt(["burger"])
    step1_section = prompt[prompt.index("STEP 1"):prompt.index("STEP 2")]
    assert "name" in step1_section.lower()


def test_prompt_step2_category_signal():
    """STEP 2 must reference the category column."""
    prompt = mc.build_system_prompt(["burger"])
    step2_section = prompt[prompt.index("STEP 2"):prompt.index("STEP 3")]
    assert "category" in step2_section.lower()


def test_prompt_step3_description_signal():
    """STEP 3 must reference the description column."""
    prompt = mc.build_system_prompt(["burger"])
    step3_section = prompt[prompt.index("STEP 3"):prompt.index("FALLBACK")]
    assert "description" in step3_section.lower()


def test_prompt_fallback_uses_others():
    """FALLBACK section must explicitly say to use Others."""
    prompt = mc.build_system_prompt(["burger"])
    fallback_section = prompt[prompt.index("FALLBACK"):]
    assert "Others" in fallback_section


def test_prompt_numbered_list():
    terms = ["grill", "salad", "juice"]
    prompt = mc.build_system_prompt(terms)
    assert "1. grill" in prompt
    assert "2. salad" in prompt
    assert "3. juice" in prompt


def test_prompt_example_lines_present():
    prompt = mc.build_system_prompt(["burger", "juice"])
    assert "1|burger" in prompt
    assert "2|juice" in prompt


def test_prompt_with_real_102_terms():
    """Smoke test: standard_term_table.csv is valid and prompt builds cleanly."""
    import polars as pl
    csv = Path(__file__).parent / "standard_term_table.csv"
    if not csv.exists():
        pytest.skip("standard_term_table.csv not present")
    df    = pl.read_csv(csv, infer_schema_length=0)
    terms = df["Standard Term"].to_list()
    assert len(terms) == 102
    prompt = mc.build_system_prompt(terms)
    for t in terms:
        assert t.strip() in prompt


# ══════════════════════════════════════════════════════════════════════════════
# 3. Checkpoint save / load
# ══════════════════════════════════════════════════════════════════════════════

def test_checkpoint_empty_when_missing(tmp_path):
    ckpt = tmp_path / "checkpoint.json"
    with patch.object(mc, "CHECKPOINT_JSON", ckpt):
        result = mc.load_checkpoint()
    assert result == {}


def test_checkpoint_roundtrip(tmp_path):
    ckpt = tmp_path / "checkpoint.json"
    data = {"0": "burger", "100": "pizza", "9999": "Others"}
    with patch.object(mc, "CHECKPOINT_JSON", ckpt):
        mc.save_checkpoint(data)
        result = mc.load_checkpoint()
    assert result == data


def test_checkpoint_no_tmp_file_left(tmp_path):
    """save_checkpoint writes via .tmp then renames — no stray .tmp stays."""
    ckpt = tmp_path / "checkpoint.json"
    with patch.object(mc, "CHECKPOINT_JSON", ckpt):
        mc.save_checkpoint({"0": "burger"})
    assert not (tmp_path / "checkpoint.tmp").exists()
    assert ckpt.exists()


def test_checkpoint_unicode_preserved(tmp_path):
    ckpt  = tmp_path / "checkpoint.json"
    data  = {"0": "mloukhieh", "1": "machboos", "2": "مندي"}
    with patch.object(mc, "CHECKPOINT_JSON", ckpt):
        mc.save_checkpoint(data)
        result = mc.load_checkpoint()
    assert result == data


def test_checkpoint_overwrites_previous(tmp_path):
    ckpt = tmp_path / "checkpoint.json"
    with patch.object(mc, "CHECKPOINT_JSON", ckpt):
        mc.save_checkpoint({"0": "burger"})
        mc.save_checkpoint({"0": "pizza"})   # overwrite
        result = mc.load_checkpoint()
    assert result["0"] == "pizza"


# ══════════════════════════════════════════════════════════════════════════════
# 4. load_input (Excel parsing)
# ══════════════════════════════════════════════════════════════════════════════

def test_load_input_basic(tmp_path):
    xlsx = tmp_path / "menu.xlsx"
    _make_xlsx(xlsx, [
        ["Chicken Shawarma", "Main", "Grilled chicken in flatbread", "shawarma"],
        ["Kunafa",           "Dessert", "Sweet cheese pastry",       "kunafa"],
    ])
    with patch.object(mc, "INPUT_XLSX", xlsx):
        rows = mc.load_input()
    assert len(rows) == 2
    assert rows[0] == (0, "Chicken Shawarma", "Main", "Grilled chicken in flatbread", "shawarma")
    assert rows[1] == (1, "Kunafa", "Dessert", "Sweet cheese pastry", "kunafa")


def test_load_input_none_cells_become_empty_string(tmp_path):
    xlsx = tmp_path / "menu.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Item Name", "Category", "Description", "existing_std_term"])
    ws.append([None, None, None, None])
    wb.save(xlsx)
    with patch.object(mc, "INPUT_XLSX", xlsx):
        rows = mc.load_input()
    assert rows[0] == (0, "", "", "", "")


def test_load_input_row_index_zero_based(tmp_path):
    xlsx = tmp_path / "menu.xlsx"
    _make_xlsx(xlsx, [["A", "B", "C", "D"], ["E", "F", "G", "H"]])
    with patch.object(mc, "INPUT_XLSX", xlsx):
        rows = mc.load_input()
    assert rows[0][0] == 0
    assert rows[1][0] == 1


def test_load_input_strips_whitespace(tmp_path):
    xlsx = tmp_path / "menu.xlsx"
    _make_xlsx(xlsx, [["  Burger  ", "  Main  ", "  Beef patty  ", "  burger  "]])
    with patch.object(mc, "INPUT_XLSX", xlsx):
        rows = mc.load_input()
    assert rows[0][1] == "Burger"
    assert rows[0][2] == "Main"
    assert rows[0][3] == "Beef patty"


def test_load_input_large_count(tmp_path):
    """100 rows all load correctly with sequential 0-based indices."""
    xlsx = tmp_path / "menu.xlsx"
    data = [[f"Item {i}", f"Cat {i}", f"Desc {i}", f"term {i}"] for i in range(100)]
    _make_xlsx(xlsx, data)
    with patch.object(mc, "INPUT_XLSX", xlsx):
        rows = mc.load_input()
    assert len(rows) == 100
    assert rows[0][0] == 0
    assert rows[99][0] == 99


# ══════════════════════════════════════════════════════════════════════════════
# 5. API response parsing (make_classifier inner _call)
# ══════════════════════════════════════════════════════════════════════════════

@pytest.mark.asyncio
async def test_parse_valid_single_item():
    fn = _make_classify_fn("0|shawarma", {"shawarma", "Others"})
    result = await fn([(0, "Shawarma Plate", "Main", "")])
    assert result[0] == "shawarma"


@pytest.mark.asyncio
async def test_parse_multiple_items():
    fn = _make_classify_fn(
        "0|shawarma\n1|burger\n2|juice",
        {"shawarma", "burger", "juice", "Others"},
    )
    batch = [(0, "A", "B", ""), (1, "C", "D", ""), (2, "E", "F", "")]
    result = await fn(batch)
    assert result == {0: "shawarma", 1: "burger", 2: "juice"}


@pytest.mark.asyncio
async def test_parse_invalid_term_becomes_others():
    fn = _make_classify_fn("0|flying_saucer_deluxe", {"burger", "Others"})
    result = await fn([(0, "Weird Item", "Cat", "")])
    assert result[0] == "Others"


@pytest.mark.asyncio
async def test_parse_case_insensitive_term_matching():
    """API returns 'Grill' (title case); valid set has 'grill' (lower)."""
    fn = _make_classify_fn("0|Grill", {"grill", "Others"})
    result = await fn([(0, "BBQ Platter", "Main", "")])
    assert result[0] == "grill"   # corrected to valid spelling


@pytest.mark.asyncio
async def test_parse_skips_lines_without_pipe():
    fn = _make_classify_fn(
        "Here are results:\n0|burger\n\nDone.",
        {"burger", "Others"},
    )
    result = await fn([(0, "Burger", "Main", "")])
    assert result[0] == "burger"


@pytest.mark.asyncio
async def test_parse_whitespace_in_response():
    fn = _make_classify_fn("  0  |  shawarma  ", {"shawarma", "Others"})
    result = await fn([(0, "Shawarma", "Main", "")])
    assert result[0] == "shawarma"


@pytest.mark.asyncio
async def test_parse_empty_response_returns_empty_dict():
    """Empty API response → caller will fill missing rows with 'Others'."""
    fn = _make_classify_fn("", {"burger", "Others"})
    result = await fn([(0, "Item", "Cat", "")])
    assert result == {}


@pytest.mark.asyncio
async def test_parse_others_passthrough():
    fn = _make_classify_fn("0|Others", {"burger", "Others"})
    result = await fn([(0, "Unclassifiable Thing", "Misc", "")])
    assert result[0] == "Others"


@pytest.mark.asyncio
async def test_description_truncated_before_sending():
    """Description > DESC_MAXLEN chars must be cut before going into the payload."""
    long_desc = "Z" * 500
    captured   = []

    async def spy_create(**kwargs):
        captured.append(kwargs["messages"][1]["content"])
        resp = MagicMock()
        resp.choices = [MagicMock()]
        resp.choices[0].message.content = "0|burger"
        return resp

    client = AsyncMock()
    client.chat.completions.create = spy_create
    bucket = mc.AsyncTokenBucket(rpm=100_000)
    sem    = asyncio.Semaphore(10)
    fn     = mc.make_classifier(client, bucket, sem, "prompt", {"burger", "Others"})

    await fn([(0, "Item", "Cat", long_desc)])

    payload = captured[0]
    # The description is the 4th pipe-delimited field: INDEX|name|cat|desc
    desc_in_payload = payload.split("|")[3]
    assert len(desc_in_payload) <= mc.DESC_MAXLEN


# ══════════════════════════════════════════════════════════════════════════════
# 6. Retry on RateLimitError
# ══════════════════════════════════════════════════════════════════════════════

@pytest.mark.asyncio
async def test_retry_on_rate_limit_then_succeed(monkeypatch):
    """First call raises RateLimitError; second call succeeds. tenacity retries."""
    from openai import RateLimitError

    # Patch asyncio.sleep so tenacity's wait doesn't actually sleep
    monkeypatch.setattr(asyncio, "sleep", AsyncMock(return_value=None))

    good_resp = MagicMock()
    good_resp.choices = [MagicMock()]
    good_resp.choices[0].message.content = "0|burger"

    call_count = 0

    async def flaky(**kwargs):
        nonlocal call_count
        call_count += 1
        if call_count == 1:
            req = httpx.Request("POST", "https://api.openai.com/v1/chat/completions")
            raise RateLimitError(
                "rate limit exceeded",
                response=httpx.Response(429, request=req, content=b"{}"),
                body={"error": {"message": "rate limit"}},
            )
        return good_resp

    client = AsyncMock()
    client.chat.completions.create = flaky
    bucket = mc.AsyncTokenBucket(rpm=100_000)
    sem    = asyncio.Semaphore(10)
    fn     = mc.make_classifier(client, bucket, sem, "prompt", {"burger", "Others"})

    result = await fn([(0, "Burger", "Main", "")])
    assert call_count == 2,  f"Expected 2 calls (1 fail + 1 success), got {call_count}"
    assert result[0] == "burger"


@pytest.mark.asyncio
async def test_retry_exhausted_raises():
    """After all 6 attempts fail the original exception re-raises (reraise=True)."""
    from openai import RateLimitError

    req = httpx.Request("POST", "https://api.openai.com/v1/chat/completions")
    call_count = 0

    async def always_fail(**kwargs):
        nonlocal call_count
        call_count += 1
        raise RateLimitError(
            "rate limit",
            response=httpx.Response(429, request=req, content=b"{}"),
            body={"error": {"message": "rate limit"}},
        )

    client = AsyncMock()
    client.chat.completions.create = always_fail
    bucket = mc.AsyncTokenBucket(rpm=100_000)
    sem    = asyncio.Semaphore(10)
    fn     = mc.make_classifier(client, bucket, sem, "prompt", {"burger", "Others"})

    # Patch asyncio.sleep at the module level so tenacity waits are instant
    with patch("asyncio.sleep", AsyncMock(return_value=None)):
        # reraise=True means the original RateLimitError surfaces (not RetryError)
        with pytest.raises(RateLimitError):
            await fn([(0, "Burger", "Main", "")])

    assert call_count == 6, f"Expected 6 attempts, got {call_count}"


# ══════════════════════════════════════════════════════════════════════════════
# 7. write_excel
# ══════════════════════════════════════════════════════════════════════════════

def test_write_excel_creates_file(tmp_path):
    rows    = [(0, "Burger", "Main", "Beef patty", "burger")]
    results = {"0": "burger"}
    out     = tmp_path / "out.xlsx"
    with patch.object(mc, "OUTPUT_XLSX", out):
        mc.write_excel(rows, results)
    assert out.exists()


def test_write_excel_five_columns_with_header(tmp_path):
    rows    = [(0, "Shawarma", "Main", "Grilled wrap", "shawarma")]
    results = {"0": "shawarma"}
    out     = tmp_path / "out.xlsx"
    with patch.object(mc, "OUTPUT_XLSX", out):
        mc.write_excel(rows, results)
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, 6)]
    assert headers[0] == "Item Name"
    assert headers[1] == "Category"
    assert headers[2] == "Description"
    assert headers[3] == "final_accurate_std_term"
    assert headers[4] == "new_std_term"


def test_write_excel_col_e_has_new_term(tmp_path):
    rows    = [(0, "Kunafa", "Dessert", "Sweet pastry", "old_term")]
    results = {"0": "kunafa"}
    out     = tmp_path / "out.xlsx"
    with patch.object(mc, "OUTPUT_XLSX", out):
        mc.write_excel(rows, results)
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    assert ws.cell(2, 5).value == "kunafa"    # col E
    assert ws.cell(2, 4).value == "old_term"  # col D preserved


def test_write_excel_missing_result_writes_empty(tmp_path):
    rows    = [(0, "Unknown", "Cat", "desc", "old")]
    results = {}   # no classification
    out     = tmp_path / "out.xlsx"
    with patch.object(mc, "OUTPUT_XLSX", out):
        mc.write_excel(rows, results)
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    # xlsxwriter writes "" as a blank cell; openpyxl reads blank cells as None
    assert ws.cell(2, 5).value in (None, "")


def test_write_excel_correct_row_count(tmp_path):
    rows    = [(i, f"Item {i}", "Cat", "desc", "term") for i in range(5)]
    results = {str(i): "burger" for i in range(5)}
    out     = tmp_path / "out.xlsx"
    with patch.object(mc, "OUTPUT_XLSX", out):
        mc.write_excel(rows, results)
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    # 1 header + 5 data rows
    assert ws.max_row == 6


def test_write_excel_multiple_rows_all_terms(tmp_path):
    rows = [
        (0, "Shawarma", "Main",    "Grilled wrap",    "old1"),
        (1, "Kunafa",   "Dessert", "Sweet pastry",    "old2"),
        (2, "Juice",    "Drinks",  "Fresh OJ",        "old3"),
    ]
    results = {"0": "shawarma", "1": "kunafa", "2": "juice"}
    out = tmp_path / "out.xlsx"
    with patch.object(mc, "OUTPUT_XLSX", out):
        mc.write_excel(rows, results)
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    assert ws.cell(2, 5).value == "shawarma"
    assert ws.cell(3, 5).value == "kunafa"
    assert ws.cell(4, 5).value == "juice"


# ══════════════════════════════════════════════════════════════════════════════
# 8. write_json
# ══════════════════════════════════════════════════════════════════════════════

def test_write_json_creates_valid_json(tmp_path):
    rows    = [(0, "Burger", "Main", "Beef patty", "burger")]
    results = {"0": "burger"}
    out     = tmp_path / "out.json"
    with patch.object(mc, "OUTPUT_JSON", out):
        mc.write_json(rows, results)
    assert out.exists()
    with open(out, encoding="utf-8") as f:
        data = json.load(f)
    assert isinstance(data, list)


def test_write_json_correct_keys(tmp_path):
    rows    = [(0, "Burger", "Main", "Beef patty", "burger")]
    results = {"0": "burger"}
    out     = tmp_path / "out.json"
    with patch.object(mc, "OUTPUT_JSON", out):
        mc.write_json(rows, results)
    with open(out, encoding="utf-8") as f:
        data = json.load(f)
    rec = data[0]
    assert set(rec.keys()) == {
        "item_name", "category", "description", "existing_std_term", "new_std_term"
    }


def test_write_json_values_correct(tmp_path):
    rows    = [(0, "Kunafa", "Dessert", "Sweet cheese pastry", "old")]
    results = {"0": "kunafa"}
    out     = tmp_path / "out.json"
    with patch.object(mc, "OUTPUT_JSON", out):
        mc.write_json(rows, results)
    with open(out, encoding="utf-8") as f:
        data = json.load(f)
    assert data[0]["item_name"]         == "Kunafa"
    assert data[0]["existing_std_term"] == "old"
    assert data[0]["new_std_term"]      == "kunafa"


def test_write_json_missing_result_empty_string(tmp_path):
    rows    = [(0, "Item", "Cat", "desc", "old")]
    results = {}
    out     = tmp_path / "out.json"
    with patch.object(mc, "OUTPUT_JSON", out):
        mc.write_json(rows, results)
    with open(out, encoding="utf-8") as f:
        data = json.load(f)
    assert data[0]["new_std_term"] == ""


def test_write_json_unicode_preserved(tmp_path):
    rows    = [(0, "مندي", "Main", "اللحم الهش", "mandi")]
    results = {"0": "mandi"}
    out     = tmp_path / "out.json"
    with patch.object(mc, "OUTPUT_JSON", out):
        mc.write_json(rows, results)
    with open(out, encoding="utf-8") as f:
        data = json.load(f)
    assert data[0]["item_name"] == "مندي"
    assert data[0]["new_std_term"] == "mandi"


def test_write_json_row_count(tmp_path):
    rows    = [(i, f"Item {i}", "Cat", "desc", "old") for i in range(10)]
    results = {str(i): "burger" for i in range(10)}
    out     = tmp_path / "out.json"
    with patch.object(mc, "OUTPUT_JSON", out):
        mc.write_json(rows, results)
    with open(out, encoding="utf-8") as f:
        data = json.load(f)
    assert len(data) == 10


# ══════════════════════════════════════════════════════════════════════════════
# 9. Integration: classify_fn + process_batch logic
# ══════════════════════════════════════════════════════════════════════════════

@pytest.mark.asyncio
async def test_batch_of_50_items_fully_classified():
    """50-item batch → all 50 indices classified."""
    valid = {"burger", "pizza", "salad", "juice", "shawarma", "Others"}
    term_cycle = ["burger", "pizza", "salad", "juice", "shawarma"] * 10

    resp_lines = "\n".join(f"{i}|{term_cycle[i]}" for i in range(50))
    fn = _make_classify_fn(resp_lines, valid)

    batch = [(i, f"Item {i}", "Cat", "") for i in range(50)]
    result = await fn(batch)

    assert len(result) == 50
    assert result[0]  == "burger"
    assert result[49] == "shawarma"


@pytest.mark.asyncio
async def test_missing_rows_in_response():
    """API only returns 2 of 3 rows → missing row absent from result (caller fills)."""
    valid = {"burger", "pizza", "Others"}
    fn = _make_classify_fn("0|burger\n2|pizza", valid)

    batch = [(0, "A", "B", ""), (1, "C", "D", ""), (2, "E", "F", "")]
    result = await fn(batch)

    assert result[0] == "burger"
    assert result[2] == "pizza"
    assert 1 not in result   # caller fills with "Others"


@pytest.mark.asyncio
async def test_classify_fn_sends_correct_item_count():
    """Verify payload has exactly N lines for N items."""
    payloads = []

    async def capture(**kwargs):
        payloads.append(kwargs["messages"][1]["content"])
        resp = MagicMock()
        resp.choices = [MagicMock()]
        resp.choices[0].message.content = "\n".join(f"{i}|burger" for i in range(5))
        return resp

    client = AsyncMock()
    client.chat.completions.create = capture
    bucket = mc.AsyncTokenBucket(rpm=100_000)
    sem    = asyncio.Semaphore(10)
    fn     = mc.make_classifier(client, bucket, sem, "prompt", {"burger", "Others"})

    await fn([(i, f"Item {i}", "Cat", "") for i in range(5)])

    assert len(payloads) == 1
    lines = [l for l in payloads[0].splitlines() if l.strip()]
    assert len(lines) == 5


@pytest.mark.asyncio
async def test_concurrent_batches_no_race_condition():
    """Run 10 batches concurrently; all should complete without interference."""
    valid = {"burger", "pizza", "Others"}
    fn    = _make_classify_fn("0|burger", valid)

    results = await asyncio.gather(*[
        fn([(j * 10 + i, f"Item {i}", "Cat", "") for i in range(5)])
        for j in range(10)
    ])

    # Each batch returned a non-empty dict
    for r in results:
        assert isinstance(r, dict)
        assert len(r) > 0
