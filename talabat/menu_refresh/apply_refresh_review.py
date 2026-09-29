"""Apply Sagar's manual review to a refresh cohort's NEW menu items.

Input   <data>/finalised.xlsx - every kNN-mapped item_key of the cohort
        (24,970 for refresh_202609), each with corrected_std_term and
        corrected_ingredients.
Output  <data>/menu_items_final.jsonl - every ADDED row (90,015): the reviewed
        kNN rows relabelled, prior_delivered/exact rows untouched.

SAME RULES AS september/apply_review_verdicts.py:
  * joined on review_id, never on the sheet's text columns
  * every corrected std_term must already be in what Tech holds, in Tech's
    casing, and resolve to a taxonomy - otherwise NOTHING is written
  * taxonomy is RE-DERIVED from the corrected std_term (the 9-value contract),
    never copied from a sheet
  * ingredients conform exact -> recover -> drop and come out in the 820-term
    file's own lowercase spelling. That file is only ever read, never grown.
  * "not applicable - non-menu item" empties the ingredient list

WHAT IS DIFFERENT, AND WHY:
  * review_id is RECOMPUTED from the rows being labelled ("MK-" +
    sha256(item_key)[:12]), so the join depends on no intermediate file. A hash
    collision refuses to write.
  * EVERY kNN key must carry a verdict. Sagar: "i want accuracy rather than
    guess work" - an unreviewed kNN row would ship a guess, so a coverage gap
    refuses to write. On refresh_202609 the reviewer kept only 62.1% of the
    labels the accept rule would have auto-shipped: 4,026 wrong labels, where
    the holdout predicted ~660.
  * the sheet's context columns are NEVER read. Excel damaged 141 of them on the
    round trip: 136 descriptions had line breaks re-encoded as _x000D_, 3
    descriptions beginning "-" were parsed as formulas and blanked, 2 item names
    became dates ("8-12" -> 2026-12-08, "Mar-50" -> 1950-03-01).
  * recover() comes from conform_to_taxonomy_file, not conform_ingredients as in
    September. conform_ingredients imports psycopg2 (Postgres is parked), and
    its rule differs by one step - it does not strip a stray leading quote, so
    '"Basmati Rice' drops instead of recovering. conform_to_taxonomy_file is
    what already conformed this cohort's untouched rows: one rule for all rows.
  * description is joined back from the scrape where a row carries "". Rows
    built before build_added_items_input.py learned to join it all have "".
  * the sheet has no Final Taxonomy column, so dish_family is carried as "" -
    parity with September's row shape, nothing invented.
  * <data>/non_menu_decisions.csv, when present, re-checks non-menu verdicts:
    it may only overturn a non-menu verdict, and every reclassification must
    land on an existing Tech std_term with ingredients from the 820-term file.
    The audit records verdict_source = sheet | non_menu_recheck.
  * taxonomy is re-derived from std_term on EVERY row, untouched ones included,
    and gated: menu/final_std_terms_taxonomy.csv once sent the std_term
    "Marketing/Non-Standard Menu" to core_food (fixed Sept 2026), and a derived
    field that is not re-derived keeps the old answer.

    python apply_refresh_review.py --dry-run
    python apply_refresh_review.py
"""
import argparse
import csv
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
import orjson

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(ROOT / "August_menu"))
from conform_to_taxonomy_file import recover, split_ings             # noqa: E402
from map_std_terms import load_taxonomy, resolve_taxonomy            # noqa: E402

TAXO = ROOT / "August_menu" / "ingredients_export_02.json"
NON_MENU = "not applicable - non-menu item"
NON_MENU_TERMS = {"marketing/non-standard menu", "service and packaging"}
HOLDOUT = {"5/5": 97.5, "4/5": 84.0, "3/5": 64.7, "2/5": 38.0, "1/5": 33.7}

# A RANKING AID for the non-menu verdicts, never a decision. September's review
# pulled 25 real drinks into non-menu through a flower/rose/petal keyword sweep,
# and their menu SECTION gave them away ("Cold Drinks", "Craft Teas"). So a
# verdict is surfaced when consumable wording appears and no UNAMBIGUOUS product
# word does. rose/flower/petal/floral/blossom are deliberately NOT product words:
# they are exactly the words that produced September's false verdicts.
CONSUMABLE = re.compile(
    r"\b(drinks?|juices?|coffee|teas?|lattes?|mocktails?|smoothies?|shakes?|"
    r"soda|water|milk|refresh\w*|tast\w*|flavou?r\w*|sips?|iced|brew\w*|"
    r"espresso|lemonade|frapp\w*|cakes?|cookies?|desserts?|sweets?|laddoo|"
    r"macarons?|pastr\w*|sandwich\w*|burgers?|pizzas?|meals?|salads?)\b", re.I)
PRODUCT = re.compile(
    r"\b(candles?|wax|crowns?|headbands?|bouquets?|hampers?|vases?|balloons?|"
    r"toys?|accessor\w*|mugs?|tumblers?|perfumes?|soaps?|plush|teddy|"
    r"keychains?|stickers?|greeting cards?|straws?|toppers?|hats?)\b", re.I)

sys.stdout.reconfigure(encoding="utf-8")


def review_id(item_key: str) -> str:
    """Identical to export_refresh_review.review_id - the id the sheet carries."""
    return "MK-" + hashlib.sha256(item_key.encode("utf-8")).hexdigest()[:12]


def load_verdicts(xlsx):
    """review_id -> (std_term, [ingredients]). Only the id and the two verdict
    columns are read; every other column went through Excel and is untrusted."""
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    rows = wb[wb.sheetnames[0]].iter_rows(values_only=True)
    head = [str(c).strip() if c is not None else "" for c in next(rows)]
    ix = {h: i for i, h in enumerate(head)}
    for need in ("review_id", "corrected_std_term", "corrected_ingredients"):
        if need not in ix:
            raise SystemExit(f"column missing from the sheet: {need!r}")
    out, dup = {}, 0
    for r in rows:
        rid = r[ix["review_id"]]
        if not rid:
            continue
        rid = str(rid).strip()
        dup += rid in out
        out[rid] = (str(r[ix["corrected_std_term"]] or "").strip(),
                    split_ings(r[ix["corrected_ingredients"]]))
    wb.close()
    return out, dup


def conform(vals, vocab, dropped):
    keep, seen = [], set()
    for v in vals:
        got = vocab.get(v.lower()) or recover(v, vocab)
        if not got:
            dropped[v.lower()] += 1
            continue
        if got not in seen:
            seen.add(got)
            keep.append(got)
    return keep


def main(args):
    D = Path(args.data)
    if not D.is_absolute():
        D = HERE / D
    src = D / "menu_items_with_std_terms.jsonl"
    xlsx = Path(args.xlsx) if args.xlsx else D / "finalised.xlsx"
    prior_f = D / "prior_item_labels.json"
    scrape = D.parent / "scrape" / "menu_items.jsonl"
    out_f = D / "menu_items_final.jsonl"
    audit_f = D / "review_verdicts_applied.csv"
    nonmenu_f = D / "non_menu_verdicts_to_confirm.csv"
    dropped_f = D / "ingredients_dropped_on_apply.csv"

    print("=" * 78)
    print("  APPLY MANUAL REVIEW - refresh cohort, new menu items")
    print("=" * 78)
    for p in (src, xlsx, prior_f, TAXO):
        if not p.exists():
            raise SystemExit(f"missing: {p}")

    canon = orjson.loads(prior_f.read_bytes())["canonical_casing"]   # Tech's vocabulary
    vocab = {r["ingredient_name"].strip().lower(): r["ingredient_name"].strip()
             for r in json.loads(TAXO.read_text(encoding="utf-8"))
             if r.get("ingredient_name")}
    tax, _ = load_taxonomy(set())
    print(f"  Tech std_term vocabulary : {len(canon):,} terms")
    print(f"  ingredient vocabulary    : {len(vocab):,} terms ({TAXO.name})")

    rows = [orjson.loads(l) for l in open(src, "rb") if l.strip()]
    print(f"  rows                     : {len(rows):,}")

    # ---- grain: one row per (branch, item) --------------------------------
    grain = Counter((r["branch_id"], r["item_key"]) for r in rows)
    if any(n > 1 for n in grain.values()):
        raise SystemExit(f"{sum(n > 1 for n in grain.values()):,} (branch, item_key) "
                         "pairs repeat - the row grain is not what this assumes")

    # ---- the join: review_id recomputed from the rows ----------------------
    knn_keys = {r["item_key"] for r in rows if r.get("std_term_method") == "knn"}
    key_of = defaultdict(set)
    for k in knn_keys:
        key_of[review_id(k)].add(k)
    if any(len(v) > 1 for v in key_of.values()):
        raise SystemExit("review_id hash collision - refusing to write")
    verdicts, dup = load_verdicts(xlsx)
    if dup:
        raise SystemExit(f"{dup:,} review_ids repeat in the sheet - refusing to write")
    unknown = [rid for rid in verdicts if rid not in key_of]
    uncovered = [rid for rid in key_of if rid not in verdicts]
    print(f"\n  verdicts in the sheet    : {len(verdicts):,}")
    print(f"  kNN keys to cover        : {len(key_of):,}")
    print(f"    verdict with no kNN key: {len(unknown):,}   (must be 0)")
    print(f"    kNN key with no verdict: {len(uncovered):,}   (must be 0 - it would ship a guess)")
    if unknown or uncovered:
        raise SystemExit("verdicts and kNN keys do not match 1:1 - refusing to write")
    by_key = {next(iter(key_of[rid])): v for rid, v in verdicts.items()}

    # ---- the non-menu re-check (Sagar, Sept 2026) --------------------------
    # 108 non-menu verdicts read as real food or drink: a keyword sweep had caught
    # "Double Puzzle Burger", "Plain Tissue Bread", V60 coffees with "floral
    # notes". Each was re-checked against its description and section; the
    # decisions live in non_menu_decisions.csv. A decision may ONLY overturn a
    # non-menu verdict, and a reclassification must carry its own label - every
    # std_term already in Tech's vocabulary, every ingredient in the 820-term
    # file - never the kNN guess.
    source_of = {}
    dec_f = D / "non_menu_decisions.csv"
    if dec_f.exists():
        dec = list(csv.DictReader(open(dec_f, encoding="utf-8-sig")))
        errs = []
        if len({d["review_id"] for d in dec}) != len(dec):
            errs.append("a review_id repeats in the decisions file")
        for d in dec:
            rid = d["review_id"]
            if rid not in verdicts:
                errs.append(f"{rid} is not in the sheet")
                continue
            k = next(iter(key_of[rid]))
            if [x.lower() for x in by_key[k][1]] != [NON_MENU]:
                errs.append(f"{rid} would overturn a verdict that is not non-menu")
            if d["decision"] == "reclassify":
                t, ings = d["std_term"].strip(), split_ings(d["ingredients"])
                if t.lower() not in canon or t.lower() in NON_MENU_TERMS:
                    errs.append(f"{rid} std_term {t!r} is not a Tech food term")
                if not ings or any(i.lower() not in vocab for i in ings):
                    errs.append(f"{rid} ingredients empty or outside the 820")
                by_key[k] = (t, ings)
                source_of[k] = "non_menu_recheck"
            elif d["decision"] != "keep_non_menu":
                errs.append(f"{rid} has unknown decision {d['decision']!r}")
        n_re = sum(d["decision"] == "reclassify" for d in dec)
        print(f"\n  non-menu re-check        : {len(dec):,} decisions ({dec_f.name}) -> "
              f"{n_re:,} reclassified, {len(dec) - n_re:,} kept non-menu")
        if errs:
            raise SystemExit("non-menu decisions failed validation - refusing to write:\n  "
                             + "\n  ".join(errs[:10]))

    # ---- verify the verdicts BEFORE touching a row -------------------------
    terms = Counter(v[0] for v in by_key.values())
    not_tech = sorted(t for t in terms if t.lower() not in canon)
    no_tax = sorted(t for t in terms
                    if not resolve_taxonomy(canon.get(t.lower(), t), tax))
    sentinel = sum(1 for v in by_key.values()
                   if [x.lower() for x in v[1]] == [NON_MENU])
    mixed = sum(1 for v in by_key.values()
                if NON_MENU in [x.lower() for x in v[1]] and len(v[1]) > 1)
    inconsistent = sum(
        1 for v in by_key.values()
        if ([x.lower() for x in v[1]] == [NON_MENU]) != (v[0].lower() in NON_MENU_TERMS))
    print(f"\n  corrected std_term       : {len(terms):,} distinct")
    print(f"    not in Tech vocabulary : {len(not_tech):,}   (must be 0) {not_tech[:8]}")
    print(f"    no taxonomy            : {len(no_tax):,}   (must be 0) {no_tax[:8]}")
    print(f"  non-menu verdicts        : {sentinel:,}")
    print(f"    sentinel mixed with ingredients      : {mixed:,}   (must be 0)")
    print(f"    non-menu label/ingredient disagree   : {inconsistent:,}   (must be 0)")
    if not_tech or no_tax or mixed or inconsistent:
        raise SystemExit("verdict verification failed - refusing to write")

    # ---- description from the scrape where the row carries none -----------
    need = {(r["branch_id"], r["item_key"]) for r in rows
            if not (r.get("description") or "").strip()}
    desc = {}
    if need and scrape.exists():
        for line in open(scrape, "rb"):
            if not line.strip():
                continue
            s = orjson.loads(line)
            k = (s.get("branch_id"), s.get("item_key"))
            if k in need and (s.get("description") or "").strip():
                desc.setdefault(k, s["description"].strip())

    # ---- apply ---------------------------------------------------------------
    dropped, st = Counter(), Counter()
    band = defaultdict(lambda: [0, 0])
    audit, first_row = [], {}
    for r in rows:
        k = (r["branch_id"], r["item_key"])
        if not (r.get("description") or "").strip() and k in desc:
            r["description"] = desc[k]
            st["description_filled"] += 1
        r.setdefault("dish_family", "")
        if r.get("std_term_method") != "knn":
            # taxonomy is a FUNCTION of std_term, so it is re-derived here too: a
            # fix to menu/final_std_terms_taxonomy.csv must reach the rows the
            # review never touched, or one term would carry two taxonomies
            x = resolve_taxonomy(r["std_term"], tax) or r.get("taxonomy")
            if x != r.get("taxonomy"):
                r["taxonomy"] = x
                st["untouched_taxonomy_rederived"] += 1
            r["review_applied"] = False
            st["untouched"] += 1
            continue

        term, ings = by_key[r["item_key"]]
        term = canon[term.lower()]                       # Tech's exact casing
        new_tax = resolve_taxonomy(term, tax)
        non_menu = [x.lower() for x in ings] == [NON_MENU]
        new_ings = [] if non_menu else conform(ings, vocab, dropped)
        before_t, before_x = r.get("std_term"), r.get("taxonomy")
        before_i = split_ings(r.get("ingredients"))

        if r["item_key"] not in first_row:               # per-KEY statistics
            first_row[r["item_key"]] = r
            kept = (before_t or "").lower() == term.lower()
            band[r.get("vote_agreement")][0] += kept
            band[r.get("vote_agreement")][1] += 1
            st["keys_reviewed"] += 1
            if not kept:
                st["std_term_changed"] += 1
                if (before_x or "") != new_tax:
                    st["taxonomy_class_changed"] += 1
            if not r.get("review_required"):
                st["auto_accepted_keys"] += 1
                st["auto_accepted_kept"] += kept
            if {x.lower() for x in before_i} != {x.lower() for x in new_ings}:
                st["ingredients_changed"] += 1
            st["non_menu"] += non_menu
            audit.append({
                "review_id": review_id(r["item_key"]),
                "item_key": r["item_key"], "item_name": r.get("item_name"),
                "description": r.get("description", ""),
                "category": r.get("category"),
                "vote_agreement": r.get("vote_agreement"),
                "auto_accepted_by_rule": "no" if r.get("review_required") else "yes",
                "std_term_machine": before_t, "std_term_final": term,
                "taxonomy_machine": before_x, "taxonomy_final": new_tax,
                "ingredients_machine": ", ".join(before_i),
                "ingredients_final": ", ".join(new_ings),
                "non_menu": "yes" if non_menu else "",
                "verdict_source": source_of.get(r["item_key"], "sheet"),
            })

        r.update(std_term=term, taxonomy=new_tax, ingredients=new_ings,
                 std_term_method="manual_review", std_term_confidence=1.0,
                 review_required=False, review_applied=True)
        st["rows_reviewed"] += 1

    menu_rows = Counter(r["item_key"] for r in rows)
    for a in audit:
        a["menu_rows"] = menu_rows[a["item_key"]]

    # ---- gates on the OUTPUT - never on the inputs --------------------------
    folded = defaultdict(set)
    for r in rows:
        folded[str(r["std_term"]).lower()].add(r["std_term"])
    gates = {
        "rows lost or added":                 abs(len(rows) - sum(grain.values())),
        "kNN guesses left unreviewed":        sum(r.get("std_term_method") == "knn" for r in rows),
        "rows still review_required":         sum(bool(r.get("review_required")) for r in rows),
        "rows with no taxonomy":              sum(not r.get("taxonomy") for r in rows),
        "ingredients not a list":             sum(not isinstance(r["ingredients"], list) for r in rows),
        "ingredient values outside the 820":  sum(x.lower() not in vocab for r in rows for x in r["ingredients"]),
        "ingredient values not lowercase":    sum(x != x.lower() for r in rows for x in r["ingredients"]),
        "std_term casing twins":              sum(len(v) > 1 for v in folded.values()),
        "reviewed std_term outside Tech":     sum(r["review_applied"] and r["std_term"].lower() not in canon for r in rows),
        "taxonomy not derived from std_term": sum((resolve_taxonomy(r["std_term"], tax) or r["taxonomy"]) != r["taxonomy"] for r in rows),
    }
    print("\n  OUTPUT GATES (every one must be 0)")
    for g, n in gates.items():
        print(f"     {'PASS' if n == 0 else 'FAIL'}  {g:36} {n:,}")
    outside = Counter(r["std_term"] for r in rows
                      if not r["review_applied"] and r["std_term"].lower() not in canon)
    print(f"  check: untouched rows whose std_term Tech does not hold: "
          f"{sum(outside.values()):,} rows / {len(outside)} terms {list(outside)[:6]}")

    # ---- what the review changed --------------------------------------------
    n = st["keys_reviewed"]
    print(f"\n  REVIEW APPLIED: {n:,} keys / {st['rows_reviewed']:,} rows "
          f"({st['untouched']:,} untouched rows)")
    print(f"    std_term changed            {st['std_term_changed']:>7,}  "
          f"({st['std_term_changed']/n*100:.1f}%)")
    print(f"      ...crossing a taxonomy class {st['taxonomy_class_changed']:>5,}  "
          f"(e.g. dessert -> beverage: the errors that mislead analytics)")
    print(f"    ingredients changed         {st['ingredients_changed']:>7,}  "
          f"({st['ingredients_changed']/n*100:.1f}%)")
    print(f"    non-menu, ingredients emptied {st['non_menu']:>5,}")
    print(f"    descriptions filled from scrape {st['description_filled']:>5,} rows")
    print(f"    untouched rows, taxonomy re-derived {st['untouched_taxonomy_rederived']:>3,}")
    print("\n  kNN accuracy MEASURED BY THE REVIEWER vs the holdout estimate:")
    for b in ("5/5", "4/5", "3/5", "2/5", "1/5"):
        ok, t = band[b]
        if t:
            print(f"     {b}  {ok:>6,}/{t:>6,} kept = {ok/t*100:5.1f}%   "
                  f"(holdout {HOLDOUT[b]:.1f}%)")
    aa, ak = st["auto_accepted_keys"], st["auto_accepted_kept"]
    if aa:
        print(f"     accept rule would have auto-shipped {aa:,} keys: {ak/aa*100:.1f}% right, "
              f"{aa-ak:,} WRONG")

    vals = sum(len(r["ingredients"]) for r in rows)
    empty = sum(not r["ingredients"] for r in rows)
    print(f"\n  ingredient values {vals:,} | dropped on apply "
          f"{sum(dropped.values()):,} | rows with none {empty:,} "
          f"(of which non-menu {sum(1 for r in rows if r['review_applied'] and not r['ingredients'] and r['std_term'].lower() in NON_MENU_TERMS):,})")
    tc = Counter(r["taxonomy"] for r in rows)
    print(f"  taxonomy ({len(tc)} values):  " +
          "  ".join(f"{t}={c:,}" for t, c in tc.most_common()))
    mc = Counter(r["std_term_method"] for r in rows)
    print(f"  method mix:  " + "  ".join(f"{m}={c:,}" for m, c in mc.most_common()))

    # ---- non-menu verdicts, likeliest-wrong first ----------------------------
    nm = []
    for a in audit:
        if a["non_menu"] != "yes":
            continue
        text = f"{a['item_name']} {a['description']} {a['category']}"
        cons, prod = CONSUMABLE.search(text), PRODUCT.search(text)
        a2 = {c: a[c] for c in ("review_id", "item_name", "description",
                                "category", "std_term_final", "menu_rows")}
        a2["looks_consumable"] = "yes" if cons and not prod else ""
        a2["signal"] = (cons.group(0) if cons else "") + (
            f" / product:{prod.group(0)}" if prod else "")
        # <- Sagar decides. A reversal here must come WITH a label: in September
        # a reversed verdict fell back to the machine's std_term, but in a
        # refresh cohort the machine label IS the kNN guess, and falling back to
        # it would reintroduce exactly the guesswork the review removed.
        a2["keep_as_non_menu"] = ""
        a2["corrected_std_term"] = ""
        a2["corrected_ingredients"] = ""
        nm.append(a2)
    nm.sort(key=lambda x: (x["looks_consumable"] != "yes", -x["menu_rows"]))
    flagged = [x for x in nm if x["looks_consumable"] == "yes"]
    print(f"\n  non-menu verdicts to confirm: {len(flagged):,} of {len(nm):,} "
          f"read as consumable (September reversed 25 of this kind)")
    for x in flagged[:12]:
        print(f"     {str(x['item_name'])[:30]:32} | {str(x['category'])[:22]:24} "
              f"| {str(x['description'])[:30]}")

    failed = [g for g, v in gates.items() if v]
    if failed:
        raise SystemExit(f"\n  GATES FAILED {failed} - refusing to write")
    if args.dry_run:
        print("\n  --dry-run: nothing written")
        return 0

    tmp = out_f.with_suffix(out_f.suffix + ".tmp")
    with open(tmp, "wb") as f:
        for r in rows:
            f.write(orjson.dumps(r) + b"\n")
    tmp.replace(out_f)
    cols = ["review_id", "item_name", "description", "category", "menu_rows",
            "vote_agreement", "auto_accepted_by_rule", "std_term_machine",
            "std_term_final", "taxonomy_machine", "taxonomy_final",
            "ingredients_machine", "ingredients_final", "non_menu",
            "verdict_source", "item_key"]
    with open(audit_f, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows({c: a[c] for c in cols} for a in audit)
    with open(nonmenu_f, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(nm[0].keys()))
        w.writeheader()
        w.writerows(nm)
    if dropped:
        with open(dropped_f, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(["value", "occurrences"])
            w.writerows(dropped.most_common())
    print(f"\n  -> {out_f.name}  ({len(rows):,} rows)")
    print(f"  -> {audit_f.name}  ({len(audit):,} keys)")
    print(f"  -> {nonmenu_f.name}  ({len(nm):,} verdicts, {len(flagged):,} flagged)")
    if dropped:
        print(f"  -> {dropped_f.name}")
    return 0


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--data", default="data/refresh_202609/stdterm")
    p.add_argument("--xlsx", help="verdict workbook (default <data>/finalised.xlsx)")
    p.add_argument("--dry-run", action="store_true")
    sys.exit(main(p.parse_args()))
