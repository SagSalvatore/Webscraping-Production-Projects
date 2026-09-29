import json
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path
import sys

sys.stdout.reconfigure(encoding="utf-8")

OUT_DIR = Path("brands/local/output")

# ── Load ───────────────────────────────────────────────────────────────────────
with open(OUT_DIR / "LOCAL_UAE.json", encoding="utf-8") as f:
    raw = json.load(f)

for r in raw:
    geo = r.pop("Geo_Coordinates", {}) or {}
    r["Geo_Lat"] = geo.get("lat", "")
    r["Geo_Lng"] = geo.get("lng", "")
    if isinstance(r.get("All_Categories"), list):
        r["All_Categories"] = " | ".join(r["All_Categories"])

df = pd.DataFrame(raw)
brand_counts = df.groupby("Brand_Name")["place_id"].nunique().rename("Location_Count")
df_summary = brand_counts.reset_index().sort_values("Location_Count", ascending=False)

high_brands     = set(df_summary[df_summary["Location_Count"] == 10]["Brand_Name"])
complete_brands = set(df_summary[df_summary["Location_Count"] < 10]["Brand_Name"])

df_high_sum = df_summary[df_summary["Brand_Name"].isin(high_brands)].reset_index(drop=True)
df_comp_sum = df_summary[df_summary["Brand_Name"].isin(complete_brands)].reset_index(drop=True)
df_high_det = df[df["Brand_Name"].isin(high_brands)].sort_values(["Brand_Name","Emirate","Name"]).reset_index(drop=True)
df_comp_det = df[df["Brand_Name"].isin(complete_brands)].sort_values(["Brand_Name","Emirate","Name"]).reset_index(drop=True)

DETAIL_COLS = ["Brand_Name","Name","Emirate","City","Address","Contact_No",
               "Google_Maps_URL","Geo_Lat","Geo_Lng","Category","All_Categories",
               "Rating","Review_Count","place_id"]

# ── Helpers ────────────────────────────────────────────────────────────────────
def thin_border():
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)

def hdr_cell(ws, row, col, val, bg, fg="FFFFFF"):
    c = ws.cell(row=row, column=col, value=val)
    c.font      = Font(name="Arial", bold=True, color=fg, size=10)
    c.fill      = PatternFill("solid", start_color=bg)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border    = thin_border()

def data_cell(c, alt=False):
    c.font      = Font(name="Arial", size=9)
    c.alignment = Alignment(vertical="center")
    c.border    = thin_border()
    if alt:
        c.fill = PatternFill("solid", start_color="F2F2F2")

def col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_summary(ws, df_sum, bg, note):
    # Title block
    max_col = 4
    ws.merge_cells(f"A1:{get_column_letter(max_col)}1")
    c = ws["A1"]
    c.value     = ws.title.upper()
    c.font      = Font(name="Arial", bold=True, size=13, color="FFFFFF")
    c.fill      = PatternFill("solid", start_color=bg)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells(f"A2:{get_column_letter(max_col)}2")
    c = ws["A2"]
    c.value     = note
    c.font      = Font(name="Arial", italic=True, size=9, color="595959")
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 30

    ws.merge_cells(f"A3:{get_column_letter(max_col)}3")
    c = ws["A3"]
    c.value     = (f"Total brands: {len(df_sum)}   |   "
                   f"Total outlet records: {int(df_sum['Location_Count'].sum()):,}")
    c.font      = Font(name="Arial", bold=True, size=10, color="404040")
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[3].height = 16

    # Column headers at row 5
    for ci, h in enumerate(["#", "Brand Name", "Locations Found", "Action / Remark"], 1):
        hdr_cell(ws, 5, ci, h, bg)
    col_widths(ws, [5, 46, 18, 55])
    ws.row_dimensions[5].height = 18

    for ri, row in df_sum.iterrows():
        r    = ri + 6
        alt  = (ri % 2 == 1)
        cnt  = int(row["Location_Count"])
        rem  = ("Cap hit — re-scrape with max=25 to capture all UAE outlets"
                if cnt == 10 else
                "Complete — all UAE locations likely already captured")
        vals = [ri + 1, row["Brand_Name"], cnt, rem]
        for ci, val in enumerate(vals, 1):
            c = ws.cell(row=r, column=ci, value=val)
            data_cell(c, alt)
            if ci == 3:
                c.alignment = Alignment(horizontal="center", vertical="center")
                if cnt == 10:
                    c.font = Font(name="Arial", size=9, bold=True, color="C65911")
        ws.row_dimensions[r].height = 16

    ws.freeze_panes = "A6"


def write_detail(ws, df_det, bg):
    hdrs     = ["Brand Name","Outlet Name","Emirate","City","Address",
                "Phone","Google Maps URL","Lat","Lng",
                "Category","All Categories","Rating","Reviews","Place ID"]
    keys     = ["Brand_Name","Name","Emirate","City","Address",
                "Contact_No","Google_Maps_URL","Geo_Lat","Geo_Lng",
                "Category","All_Categories","Rating","Review_Count","place_id"]
    widths   = [28, 32, 16, 14, 42, 17, 52, 10, 10, 22, 30, 8, 8, 28]

    for ci, h in enumerate(hdrs, 1):
        hdr_cell(ws, 1, ci, h, bg)
    col_widths(ws, widths)
    ws.row_dimensions[1].height = 20

    for ri, row in df_det.iterrows():
        r   = ri + 2
        alt = ri % 2 == 1
        for ci, key in enumerate(keys, 1):
            val = row[key]
            try:
                if pd.isna(val):
                    val = ""
            except Exception:
                pass
            c = ws.cell(row=r, column=ci, value=val)
            data_cell(c, alt)
            if key == "Google_Maps_URL" and val:
                c.hyperlink = str(val)
                c.font      = Font(name="Arial", size=9, color="0563C1", underline="single")
        ws.row_dimensions[r].height = 15

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(hdrs))}1"


# ── FILE 1 — HIGH POTENTIAL ────────────────────────────────────────────────────
ORANGE = "C65911"
wb1 = Workbook()
ws1a = wb1.active
ws1a.title = "High Potential Summary"
write_summary(
    ws1a, df_high_sum, ORANGE,
    "These 35 brands each returned exactly 10 results (our search cap).  They likely have MORE UAE locations undiscovered.\n"
    "Target them separately with a higher max (25-30) to get full coverage."
)
ws1b = wb1.create_sheet("All Locations")
write_detail(ws1b, df_high_det[DETAIL_COLS], ORANGE)
wb1.save(OUT_DIR / "LOCAL_UAE_HIGH_POTENTIAL.xlsx")
print(f"Saved LOCAL_UAE_HIGH_POTENTIAL.xlsx  [{len(df_high_sum)} brands, {len(df_high_det)} outlets]")

# ── FILE 2 — COMPLETE COVERAGE ─────────────────────────────────────────────────
NAVY = "17375E"
wb2 = Workbook()
ws2a = wb2.active
ws2a.title = "Complete Coverage Summary"
write_summary(
    ws2a, df_comp_sum, NAVY,
    "These 55 brands returned fewer than 10 results.  Google Maps has surfaced all known UAE locations for each — coverage is likely complete.\n"
    "No re-scraping needed unless the brand opens new outlets."
)
ws2b = wb2.create_sheet("All Locations")
write_detail(ws2b, df_comp_det[DETAIL_COLS], NAVY)
wb2.save(OUT_DIR / "LOCAL_UAE_COMPLETE_COVERAGE.xlsx")
print(f"Saved LOCAL_UAE_COMPLETE_COVERAGE.xlsx  [{len(df_comp_sum)} brands, {len(df_comp_det)} outlets]")

print("Done.")
