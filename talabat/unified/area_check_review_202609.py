"""Read-only review of the PUBLISHED talabat_unified_202609.jsonl - nothing in
the deliverable is changed (Sagar, Sept 18: "ill first send the deliverable to
tech team without changing any existing data").

Answers three questions in one workbook, review/unified_202609_area_check.xlsx:

  1 Where does the address give the SAME area name as the record's area?
    Every Talabat restaurant (location records excluded) classified by how its
    location.raw relates to location.area, with the restaurant name and its
    Talabat URL. The test is the builder's own address_names_area, which is
    strict: 'Al Barsha 3' never backs 'Al Barsha 1'. Also listed: brands whose
    branches share one area name while more than 1 km apart (the Baskin
    Robbins pattern).
  2 Is Baskin Robbins resolved? Every Baskin Robbins record, 202608 values
    beside 202609 values.
  3 Are ingredients normalised? Lowercase, trimmed, inside the 820-term
    vocabulary, no duplicates within an item.

    python area_check_review_202609.py
"""
import csv
import json
import math
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import ijson
import orjson
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(ROOT / "export"))
sys.stdout.reconfigure(encoding="utf-8")
from build_unified_202609 import (NA, TAXO, address_names_area, emirate,   # noqa: E402
                                  is_area_name_raw)
from export_to_json import parse_address_components                       # noqa: E402
from rectify_areas import AreaRectifier                                   # noqa: E402

FILE = HERE / "data" / "talabat_unified_202609.jsonl"
BASE = HERE / "data" / "talabat_unified_202608.jsonl"
OUT = HERE / "review" / "unified_202609_area_check.xlsx"
URL_SOURCES = (   # most recent first; each maps a branch id to its Talabat page
    (ROOT / "menu_refresh" / "data" / "refresh_202609" / "scrape_targets_202609.jsonl", "branch_id", "url"),
    (ROOT / "september" / "data" / "sept_menu_targets.jsonl", "branch_id", "url"),
    (ROOT / "images" / "data" / "august_logo_targets.jsonl", "source_id", "map_url"),
    (ROOT / "Restaurant Identifier" / "data" / "restaurant_urls_for_scraping.jsonl", "source_id", "map_url"),
    (ROOT / "data" / "urls" / "talabat_restaurant_urls_run2.jsonl", "branch_id", "url"),
    (ROOT / "data" / "urls" / "talabat_restaurant_urls.jsonl", "branch_id", "url"),
)
FAR_KM = 1.0

SAME_FULL = "Same area - full address names the area"
SAME_NAME = "Same area - address is just the area name"
DIFF_NAME = "Different area - address is an old area name"
DIFF_FULL = "Different area - full address names another area"
NUMBER = "Area number not confirmed by the address"
NO_AREA = "Address names no area"
EMIRATE = "Address names another emirate"
NO_ADDR = "No address (NA)"
NO_REC_AREA = "Record has no area"
ORDER = (SAME_FULL, SAME_NAME, DIFF_NAME, DIFF_FULL, NUMBER, NO_AREA, EMIRATE, NO_ADDR, NO_REC_AREA)


def km(a, b):
    la1, lo1, la2, lo2 = map(math.radians, (a[0], a[1], b[0], b[1]))
    h = (math.sin((la2 - la1) / 2) ** 2
         + math.cos(la1) * math.cos(la2) * math.sin((lo2 - lo1) / 2) ** 2)
    return 2 * 6371.0 * math.asin(math.sqrt(h))


def point(g):
    try:
        return float(g["lat"]), float(g["lng"])
    except (TypeError, ValueError, KeyError):
        return None


def urls():
    out = {}
    for path, kid, kurl in URL_SOURCES:
        if not path.exists():
            continue
        for line in open(path, "rb"):
            r = orjson.loads(line)
            u = r.get(kurl)
            if u and "talabat.com" in u:
                out.setdefault(str(r.get(kid)), u)
    return out


def classify(rect, raw, area, city):
    """-> (category, detail) for one restaurant's raw address against its area."""
    if not raw or raw == NA:
        return NO_ADDR, ""
    if not area:
        return NO_REC_AREA, ""
    single = " - " not in raw and "," not in raw
    if is_area_name_raw(raw, area):
        return SAME_NAME, ""
    pc = parse_address_components(raw)[0]
    if pc and city and emirate(pc) != emirate(city):
        return EMIRATE, f"address says {pc}, city is {city}"
    ok, how = address_names_area(rect, area, city, raw)
    if ok:
        return (SAME_NAME if single else SAME_FULL), how
    if how.startswith("address names another area"):
        return (DIFF_NAME if single else DIFF_FULL), how.split(": ", 1)[-1]
    if how.startswith("address does not confirm"):
        return NUMBER, ""
    return NO_AREA, ""


def main():
    rect = AreaRectifier()
    url = urls()
    vocab = {r["ingredient_name"].strip().lower() for r in json.loads(TAXO.read_text(encoding="utf-8"))}
    sept_ids = {str(r["source_id"]) for r in
                ijson.items(open(ROOT / "september" / "data" / "September_export.json", "rb"), "item")}

    rows, cats = [], Counter()
    groups = defaultdict(list)
    baskin = []
    ing, ing_distinct = Counter(), set()
    ing_bad = defaultdict(list)
    for line in open(FILE, "rb"):
        r = orjson.loads(line)
        sid = str(r["source_id"])
        ver = bool(r.get("is_verified_location"))
        L = r.get("location") or {}
        raw, area, city = L.get("raw"), L.get("area"), L.get("city")
        # ---- ingredients, every item of every record
        for it in r["menu_items"]:
            v = it.get("ingredients")
            ing["items"] += 1
            if not isinstance(v, list):
                ing["ingredients not a list"] += 1
                ing_bad["ingredients not a list"].append((r["name"], it["name"], repr(v)[:60]))
                continue
            if v:
                ing["items with ingredients"] += 1
            seen = set()
            for x in v:
                ing["ingredient entries"] += 1
                ing_distinct.add(x)
                problems = []
                if not isinstance(x, str):
                    problems.append("not text")
                else:
                    if x != x.lower():
                        problems.append("not lowercase")
                    if x != x.strip() or "  " in x:
                        problems.append("extra spaces")
                    if x.strip().lower() not in vocab:
                        problems.append("outside the 820-term vocabulary")
                    if x in seen:
                        problems.append("duplicate within the item")
                    seen.add(x)
                for p in problems:
                    ing[p] += 1
                    if len(ing_bad[p]) < 200:
                        ing_bad[p].append((r["name"], it["name"], x))
        if "baskin" in str(r.get("name")).lower():
            baskin.append(r)
        if ver:
            continue
        cohort = "september" if sid in sept_ids else "202608"
        cat, detail = classify(rect, raw, area, city)
        cats[cat] += 1
        rows.append((sid, cohort, r.get("name"), url.get(sid, ""), city, area, raw, cat, detail))
        pt = point(r.get("geo") or {})
        if pt and area:
            groups[(r.get("chain_id"), area)].append((sid, r.get("name"), pt, city, raw))

    ing_distinct = len(ing_distinct)

    # ---- brands whose branches share one area name while >1 km apart
    same_area = []
    for (cid, area), m in groups.items():
        if len(m) < 2:
            continue
        spread = max(km(a[2], b[2]) for i, a in enumerate(m) for b in m[i + 1:])
        if spread > FAR_KM:
            same_area.append((cid, area, m, spread))
    same_area.sort(key=lambda g: -g[3])

    # ---- Baskin Robbins: 202608 values beside 202609
    before = {}
    for line in open(BASE, "rb"):
        if b"askin" not in line:
            continue
        b = orjson.loads(line)
        if "baskin" in str(b.get("name")).lower() and not b.get("is_verified_location"):
            before[str(b["source_id"])] = b

    write(rows, cats, same_area, baskin, before, url, rect, ing, ing_bad, ing_distinct)
    print(f"restaurants classified {len(rows):,} | urls found {sum(1 for x in rows if x[3]):,}")
    for c in ORDER:
        print(f"  {cats[c]:>6,}  {c}")
    print(f"brand/area pairs >1 km apart: {len(same_area)} ({sum(len(g[2]) for g in same_area)} branches)")
    print("ingredients:", {k: v for k, v in ing.items()}, "| distinct", ing_distinct)
    print(f"-> {OUT}")


def write(rows, cats, same_area, baskin, before, url, rect, ing, ing_bad, ing_distinct):
    F = "Arial"
    head_font, body_font = Font(name=F, bold=True, color="FFFFFF"), Font(name=F, size=10)
    link_font = Font(name=F, size=10, color="0563C1", underline="single")
    head_fill = PatternFill("solid", fgColor="305496")
    same_fill = PatternFill("solid", fgColor="E2EFDA")
    wb = Workbook()

    def table(ws, header, data, widths, link_col=None, fill_rule=None):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for row in data:
            ws.append(list(row))
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for r in ws.iter_rows(min_row=2):
            fill = fill_rule(r) if fill_rule else None
            for c in r:
                c.font = body_font
                if fill:
                    c.fill = fill
            if link_col is not None:
                c = r[link_col]
                if isinstance(c.value, str) and c.value.startswith("http"):
                    c.hyperlink, c.font = c.value, link_font
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    # ------------------------------------------------------------ Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "talabat_unified_202609.jsonl - area, Baskin Robbins and ingredient checks (read-only)"
    ws["A1"].font = Font(name=F, bold=True, size=13)
    ws["A2"] = ("Source: unified/data/talabat_unified_202609.jsonl as published Sept 17 2026 "
                "(sha256 f82cfebe...0af6864). Nothing in the deliverable was changed.")
    ws["A2"].font = Font(name=F, italic=True, size=9)
    ws["A4"], ws["B4"], ws["C4"] = "How the address (location.raw) relates to the area", "Restaurants", "Share"
    for c in (ws["A4"], ws["B4"], ws["C4"]):
        c.font, c.fill = head_font, head_fill
    first = 5
    for i, cat in enumerate(ORDER):
        r = first + i
        ws.cell(r, 1, cat)
        ws.cell(r, 2, f"=COUNTIF(Raw_vs_Area!$H:$H,$A{r})")
        ws.cell(r, 3, f"=IF($B${first + len(ORDER)}=0,0,B{r}/$B${first + len(ORDER)})")
        ws.cell(r, 3).number_format = "0.0%"
        ws.cell(r, 2).number_format = "#,##0"
        if cat in (SAME_FULL, SAME_NAME):
            for c in range(1, 4):
                ws.cell(r, c).fill = same_fill
    tot = first + len(ORDER)
    ws.cell(tot, 1, "Total Talabat restaurants (Google location records excluded)")
    ws.cell(tot, 2, f"=SUM(B{first}:B{tot - 1})")
    ws.cell(tot, 2).number_format = "#,##0"
    ws.cell(tot + 1, 1, "Same area name in the address (the two green rows)")
    ws.cell(tot + 1, 2, f"=B{first}+B{first + 1}")
    ws.cell(tot + 1, 3, f"=IF($B${tot}=0,0,B{tot + 1}/$B${tot})")
    ws.cell(tot + 1, 2).number_format, ws.cell(tot + 1, 3).number_format = "#,##0", "0.0%"
    for rr in (tot, tot + 1):
        for c in range(1, 4):
            ws.cell(rr, c).font = Font(name=F, bold=True)
    notes = [
        "Counts are COUNTIF formulas over the Raw_vs_Area sheet - filter its 'category' column to see each group.",
        "'Same area' uses the builder's strict test: a numbered area needs its own number "
        "('Al Barsha 3' never backs 'Al Barsha 1'); Google spellings such as 'Al Qouz Ind.first' are read.",
        "'Different area - address is an old area name' is the group Sagar asked to leave untouched (Sept 17).",
    ]
    for i, n in enumerate(notes):
        ws.cell(tot + 3 + i, 1, n).font = Font(name=F, italic=True, size=9)

    r0 = tot + 7
    ws.cell(r0, 1, "Ingredient check (every menu item)").font = head_font
    ws.cell(r0, 1).fill = head_fill
    ws.cell(r0, 2, "Count").font = head_font
    ws.cell(r0, 2).fill = head_fill
    lines = [("Menu items checked", ing["items"]), ("Items with ingredients", ing["items with ingredients"]),
             ("Ingredient entries", ing["ingredient entries"]), ("Distinct ingredients", ing_distinct),
             ("Entries NOT lowercase", ing["not lowercase"]), ("Entries with extra spaces", ing["extra spaces"]),
             ("Entries outside the 820-term vocabulary", ing["outside the 820-term vocabulary"]),
             ("Duplicate entries within one item", ing["duplicate within the item"]),
             ("Items whose ingredients are not a list", ing["ingredients not a list"])]
    for i, (k, v) in enumerate(lines, 1):
        ws.cell(r0 + i, 1, k)
        ws.cell(r0 + i, 2, v).number_format = "#,##0"
    ws.cell(r0 + len(lines) + 1, 1,
            "Ingredient counts are computed by script over all 2,280,305 items of the published file "
            "(the items are not in this workbook).").font = Font(name=F, italic=True, size=9)
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold if c.font else False)
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 10

    # ------------------------------------------------------------ Raw_vs_Area
    order = {c: i for i, c in enumerate(ORDER)}
    rows.sort(key=lambda x: (order[x[7]], str(x[5]), str(x[2])))
    table(wb.create_sheet("Raw_vs_Area"),
          ["source_id", "cohort", "restaurant_name", "talabat_url", "city", "area", "raw_address",
           "category", "detail"],
          rows, [10, 10, 34, 60, 14, 28, 70, 44, 40], link_col=3,
          fill_rule=lambda r: same_fill if r[7].value in (SAME_FULL, SAME_NAME) else None)

    # ------------------------------------------------------------ Brand_Same_Area
    data = []
    for n, (cid, area, m, spread) in enumerate(same_area, 1):
        for sid, name, pt, city, raw in sorted(m, key=lambda x: x[1] or ""):
            data.append((n, str(cid), area, len(m), round(spread, 1), sid, name, url.get(sid, ""),
                         city, raw, round(pt[0], 6), round(pt[1], 6)))
    table(wb.create_sheet("Brand_Same_Area"),
          ["group", "chain_id", "shared_area", "branches_in_group", "max_distance_km", "source_id",
           "restaurant_name", "talabat_url", "city", "raw_address", "lat", "lng"],
          data, [7, 20, 28, 10, 11, 10, 34, 60, 12, 60, 11, 11], link_col=7)

    # ------------------------------------------------------------ Baskin_Robbins
    data = []
    for r in sorted(baskin, key=lambda x: (bool(x.get("is_verified_location")), str(x["source_id"]))):
        sid = str(r["source_id"])
        L = r["location"]
        ver = bool(r.get("is_verified_location"))
        b = before.get(sid) if not ver else None
        bl = (b or {}).get("location") or {}
        cat, _ = classify(rect, L.get("raw"), L.get("area"), L.get("city")) if not ver else ("", "")
        data.append((sid, "Google location record" if ver else "Talabat restaurant", r["name"],
                     url.get(sid, "") if not ver else "", bl.get("city", ""), bl.get("area", ""),
                     bl.get("raw", ""), L.get("city"), L.get("area"), L.get("raw"), cat,
                     r.get("contact_phone"), r.get("maps_url")))
    table(wb.create_sheet("Baskin_Robbins"),
          ["source_id", "record", "restaurant_name", "talabat_url", "city_202608", "area_202608",
           "raw_202608", "city_202609", "area_202609", "raw_202609", "raw_vs_area_202609",
           "contact_phone", "maps_url"],
          data, [10, 20, 18, 58, 12, 22, 50, 12, 22, 50, 40, 18, 50], link_col=3)

    # ------------------------------------------------------------ Ingredients
    data = [(p, rn, itn, x) for p, lst in ing_bad.items() for rn, itn, x in lst]
    table(wb.create_sheet("Ingredient_Issues"),
          ["problem", "restaurant_name", "item_name", "ingredient"],
          data or [("none found", "", "", "")], [34, 34, 40, 40])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)


if __name__ == "__main__":
    main()
